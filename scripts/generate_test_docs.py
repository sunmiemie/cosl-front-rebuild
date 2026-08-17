# -*- coding: utf-8 -*-
"""生成 COSL 测试管理系统测试用例 Excel 与 XMind 思维导图"""

import json
import uuid
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
from xmind_test_tree import ROOT_CHILDREN

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "docs" / "test"

# 测试用例数据：(编号, 模块, 优先级, 名称, 前置条件, 步骤, 测试数据, 预期结果, 类型, 自动化, 备注)
TEST_CASES = [
    # ===== P0 冒烟 =====
    ("S-01", "认证", "P0", "账号密码登录", "系统可访问，测试账号已创建",
     "1. 打开 /login\n2. 输入用户名密码\n3. 点击登录\n4. 使用错误密码重试",
     "正确账号: test_user / Test@123\n错误密码: wrong_pass",
     "正确账号跳转 /dashboard，Token 写入 Cookie；错误账号提示失败且不跳转",
     "冒烟", "建议", ""),
    ("S-02", "权限", "P0", "动态路由权限过滤", "已准备多角色账号",
     "1. 分别用各角色登录\n2. 查看左侧菜单\n3. 直接访问未授权 URL",
     "工单员(role11)、审核员(role16)、普通用户",
     "菜单仅显示授权项；未授权路由跳转 403",
     "冒烟", "建议", ""),
    ("S-03", "首页", "P0", "待处理任务中心", "存在待审工艺或报告",
     "1. 登录进入 /dashboard\n2. 查看任务分组 Tab\n3. 点击任务跳转",
     "有待审数据的测试环境",
     "任务按业务分组展示；点击正确跳转；30s 轮询刷新",
     "冒烟", "否", ""),
    ("S-04", "测试工艺", "P0", "工艺全流程", "有板型数据；账号有 TEST_PROCESS_CREATE 和 TEST_PROCESS_REVIEW",
     "1. 创建测试工艺并保存\n2. 提交审核\n3. 审核员通过审核\n4. 确认版本激活",
     "选择已有板型；功能性或目检工艺",
     "工艺状态变为已激活；工单创建时可选该工艺",
     "冒烟", "建议", "核心链路"),
    ("S-05", "测试工单", "P0", "生产工单发布执行", "有激活工艺、批次、板卡；测试客户端可用",
     "1. 创建生产测试工单\n2. 发布工单\n3. 客户端领取执行\n4. 确认报告生成",
     "生产批次+板型+已激活工艺",
     "工单状态流转正确；报告成功生成",
     "冒烟", "建议", "依赖 cosl-operator-tauri"),
    ("S-06", "报告审核", "P0", "报告审核通过", "有待审报告；审核员账号",
     "1. 进入报告审核列表\n2. 打开待审报告\n3. 点击通过",
     "普通生产测试报告",
     "审核成功；PDF 可下载；dashboard 待办消失",
     "冒烟", "建议", ""),
    ("S-07", "报告审核", "P0", "报告审核打回", "有待审报告",
     "1. 打开待审报告\n2. 填写打回意见\n3. 确认打回",
     "打回意见: 测试数据不完整",
     "报告状态回退；可重新测试并再次提交",
     "冒烟", "建议", ""),

    # ===== P1 认证与权限 =====
    ("A-01", "认证", "P1", "SSO 登录", "SSO 服务可用",
     "1. 从 SSO 入口跳转\n2. 完成 SSO 认证\n3. 回调带 token 参数",
     "SSO 回调 URL: ?token=xxx",
     "Token 正确写入；自动跳转 dashboard",
     "功能", "否", ""),
    ("A-02", "认证", "P1", "登出", "已登录状态",
     "1. 点击用户菜单登出\n2. 尝试访问业务页面",
     "—", "Cookie 清除；重定向登录页",
     "功能", "建议", ""),
    ("A-03", "权限", "P1", "按钮级权限控制", "准备不同权限账号",
     "1. 进入电路板型号页\n2. 用 role7 账号操作\n3. 用无权限账号操作",
     "role7 有新增/编辑权限",
     "有权限显示按钮且可操作；无权限按钮隐藏或禁用",
     "权限", "建议", ""),
    ("A-04", "权限", "P1", "角色权限分配", "管理员账号",
     "1. 新增测试角色\n2. 拖拽分配权限\n3. 给用户分配新角色\n4. 该用户重新登录验证",
     "分配 TEST_ORDER_CREATE_BASIC",
     "用户菜单和操作与分配权限一致",
     "权限", "否", ""),

    # ===== P1 装备管理 =====
    ("B-01", "装备管理", "P1", "电路板树形浏览", "有产品/仪器/板型/批次数据",
     "1. 进入电路板管理\n2. 展开树形节点\n3. 点击批次查看详情",
     "生产 Tab",
     "层级正确：产品→仪器→板型→批次；详情数据完整",
     "功能", "否", ""),
    ("B-02", "装备管理", "P1", "生产/临时板卡切换", "两类数据均存在",
     "1. 切换生产/临时 Tab\n2. 对比列表数据",
     "—", "Tab 数据隔离；互不影响",
     "功能", "否", ""),
    ("B-03", "装备管理", "P1", "二维码生成", "有电路板记录",
     "1. 选择板卡\n2. 生成二维码\n3. 扫码验证内容",
     "—", "二维码可生成；内容含正确 SN/标识",
     "功能", "否", ""),
    ("B-04", "装备管理", "P1", "其他设备 Excel 导入", "准备合法和非法 Excel",
     "1. 下载模板(如有)\n2. 导入合法文件\n3. 导入重复 SN 文件\n4. 导入格式错误文件",
     "合法 xlsx；重复 SN；空文件",
     "合法导入成功；重复/非法有明确错误提示",
     "功能", "否", ""),
    ("B-05", "装备型号", "P1", "电路板型号 CRUD", "role7 和无权限账号",
     "1. role7 新增型号\n2. role7 编辑型号\n3. 无权限账号尝试操作",
     "型号名: TEST-BOARD-001",
     "有权限 CRUD 成功；无权限不可操作",
     "权限", "否", ""),
    ("B-06", "装备型号", "P1", "电路板组管理", "有多个板型",
     "1. 新增电路板组\n2. 添加组内板型\n3. 删除板型\n4. 删除有关联的组",
     "组名: TEST-GROUP-01",
     "组内配置保存正确；有关联时删除有提示",
     "功能", "否", ""),
    ("B-07", "装备管理", "P1", "仪器管理 CRUD", "有仪器型号",
     "1. 新增仪器\n2. 编辑仪器\n3. 删除仪器\n4. 条件查询",
     "SN: INST-TEST-001",
     "CRUD 正常；查询结果准确",
     "功能", "否", ""),
    ("B-08", "装备管理", "P1", "夹具管理 CRUD", "有夹具型号",
     "1. 新增夹具并绑定型号\n2. 编辑序列号\n3. 删除",
     "序列号: FIX-TEST-001",
     "绑定关系正确；CRUD 正常",
     "功能", "否", ""),
    ("B-09", "装备管理", "P1", "主机(PSI)管理 CRUD", "有主机型号",
     "1. 新增主机\n2. 编辑状态\n3. 删除",
     "SN: PSI-TEST-001",
     "CRUD 正常；状态选项正确",
     "功能", "否", ""),
    ("B-10", "装备管理", "P1", "其他设备年检更新", "有其他设备记录",
     "1. 选择设备\n2. 更新年检日期\n3. 查看状态变化",
     "年检日期: 2026-06-01",
     "年检信息保存；到期状态正确",
     "功能", "否", ""),

    # ===== P1 测试工艺 =====
    ("C-01", "测试工艺", "P1", "新增功能性测试工艺", "有板型；TEST_PROCESS_CREATE 权限",
     "1. 选板型新增工艺\n2. 选功能性测试类型\n3. 上传测试脚本\n4. 绑定设备\n5. 保存草稿",
     "脚本文件: test_script.py",
     "草稿保存成功；脚本和设备绑定正确",
     "功能", "否", ""),
    ("C-02", "测试工艺", "P1", "新增目检测试工艺", "有板型",
     "1. 选目检测试类型\n2. 配置流程图\n3. 配置扫码项\n4. 预览正反面",
     "—", "流程图保存；预览正确显示",
     "功能", "否", ""),
    ("C-03", "测试工艺", "P1", "新增 X-RAY 工艺", "有板型",
     "1. 选 X-RAY 类型\n2. 配置参数\n3. 保存",
     "—", "参数保存成功",
     "功能", "否", ""),
    ("C-04", "测试工艺", "P1", "工艺版本切换", "同一板型有多版本工艺",
     "1. 查看版本列表\n2. 切换版本\n3. 创建工单验证可用版本",
     "—", "仅激活版本可用于工单",
     "功能", "否", ""),
    ("C-05", "测试工艺", "P1", "工艺审核打回", "有待审工艺",
     "1. 审核员打开工艺\n2. 填写打回原因\n3. 提交打回\n4. 创建人重新编辑",
     "打回原因: 流程图缺失步骤",
     "打回成功；可重新编辑提交",
     "功能", "否", ""),
    ("C-06", "测试工艺", "P1", "导出工艺树 Excel", "工艺列表有数据",
     "1. 点击导出\n2. 打开 Excel 核对",
     "—", "导出文件字段与列表一致",
     "功能", "否", ""),
    ("C-07", "测试工艺", "P1", "HALT-HASS 工艺", "有 HALT-HASS 权限",
     "1. 创建 HALT-HASS 工艺\n2. 配置安装指南\n3. 预览流程图\n4. 提交审核",
     "—", "专项编辑器功能正常；审核流程通",
     "功能", "否", ""),

    # ===== P1 测试工单 =====
    ("D-01", "测试工单", "P1", "生产工单创建", "有批次/板型/激活工艺",
     "1. 进入生产工单 Tab\n2. 点击新增\n3. 填写必填项\n4. 保存",
     "批次+板型+工艺组合",
     "工单创建成功；列表可见",
     "功能", "否", ""),
    ("D-02", "测试工单", "P1", "批量发布工单", "有多条未发布工单",
     "1. 多选工单\n2. 点击批量发布\n3. 查看状态",
     "选择 3 条以上",
     "全部发布成功；状态变为已发布",
     "功能", "否", ""),
    ("D-03", "测试工单", "P1", "工单作废/删除", "有各状态工单",
     "1. 删除未发布工单\n2. 作废已发布工单\n3. 尝试删除已执行工单",
     "—", "符合状态规则；已执行不可删",
     "功能", "否", ""),
    ("D-04", "测试工单", "P1", "导出工单 JSON", "有已发布工单",
     "1. 选择工单导出\n2. 检查 JSON 格式\n3. 客户端导入验证",
     "—", "JSON 格式正确；客户端可识别",
     "功能", "否", "依赖客户端"),
    ("D-05", "测试工单", "P1", "导入报告", "有已完成工单",
     "1. 选择工单\n2. 导入报告文件\n3. 查看关联结果",
     "合法报告 JSON",
     "报告正确关联工单",
     "功能", "否", ""),
    ("D-06", "测试工单", "P1", "临时测试工单", "有临时批次/板型",
     "1. 切换临时工单 Tab\n2. 创建临时工单\n3. 发布",
     "临时批次体系",
     "独立于生产体系；流程正常",
     "功能", "否", ""),
    ("D-07", "测试工单", "P1", "HALT-HASS 工单", "TEST_ORDER_CREATE_HALT_HASS 权限",
     "1. 进入 HALT-HASS Tab\n2. 创建工单\n3. 绑定批次+板型+工艺\n4. 发布",
     "—", "专项工单创建发布正常",
     "功能", "否", ""),
    ("D-08", "测试工单", "P1", "水循环工单", "TEST_ORDER_CREATE_WATER_CYCLE 权限",
     "1. 创建水循环工单\n2. 填写专项字段\n3. 发布",
     "—", "专项字段校验和发布正常",
     "功能", "否", ""),
    ("D-09", "测试工单", "P1", "高温箱工单", "TEST_ORDER_CREATE_HIGH_TEMP_BOX 权限",
     "1. 创建高温箱工单\n2. 发布",
     "—", "创建发布正常",
     "功能", "否", ""),
    ("D-10", "测试工单", "P1", "振动台工单", "TEST_ORDER_CREATE_VIBRATION_TABLE 权限",
     "1. 创建振动台工单\n2. 发布",
     "—", "创建发布正常",
     "功能", "否", ""),
    ("D-11", "测试工单", "P1", "高温高压井工单", "TEST_ORDER_CREATE_HIGH_TEMP_PRESSURE 权限",
     "1. 创建高温高压井工单\n2. 发布",
     "—", "创建发布正常",
     "功能", "否", ""),
    ("D-12", "测试工单", "P1", "模拟井工单", "TEST_ORDER_CREATE_SIMULATED_WELL 权限",
     "1. 创建模拟井工单\n2. 发布",
     "—", "创建发布正常",
     "功能", "否", ""),

    # ===== P1 报告与审核 =====
    ("E-01", "测试报告", "P1", "报告列表多条件查询", "有多条报告",
     "1. 按板型筛选\n2. 按日期筛选\n3. 按状态筛选\n4. 分页翻页",
     "日期范围: 近30天",
     "筛选准确；分页正常",
     "功能", "否", ""),
    ("E-02", "测试报告", "P1", "生产报告预览", "有已完成报告",
     "1. 打开报告预览\n2. 查看 2D 图\n3. 查看弓曲度\n4. 查看 X-RAY 数据",
     "—", "各预览组件正确渲染",
     "功能", "否", ""),
    ("E-03", "测试报告", "P1", "PDF 生成与下载", "有审核通过报告",
     "1. 点击生成 PDF\n2. 下载 PDF\n3. 检查内容完整性",
     "—", "PDF 生成成功；含签名和测试数据",
     "功能", "否", ""),
    ("E-04", "报告审核", "P1", "普通报告审核通过", "有待审普通报告",
     "1. 进入审核列表\n2. 打开报告\n3. 审核通过",
     "—", "状态变更；可下载 PDF",
     "功能", "建议", ""),
    ("E-05", "报告审核", "P1", "功能性测试报告审核", "有待审功能性报告",
     "1. 打开功能性报告审核页\n2. 查看预览\n3. 通过/打回",
     "—", "FunctionalTestingReportPreview 正确",
     "功能", "否", ""),
    ("E-06", "报告审核", "P1", "临时报告审核", "有待审临时报告",
     "1. 打开临时报告审核\n2. 预览\n3. 审核操作",
     "—", "TempReportPreview 正确",
     "功能", "否", ""),
    ("E-07", "报告审核", "P1", "HALT-HASS 报告审核", "REPORT_REVIEW_HALT_HASS 权限",
     "1. 进入专项 Tab\n2. 查看曲线图\n3. 审核通过/打回",
     "—", "专项审核功能正常",
     "功能", "否", ""),
    ("E-08", "报告审核", "P1", "水循环报告审核", "REPORT_REVIEW_WATER_CYCLE 权限",
     "1. 查看待审列表\n2. 详情弹窗\n3. 审核操作",
     "—", "审核通过/打回正常",
     "功能", "否", ""),
    ("E-09", "报告审核", "P1", "高温箱报告审核", "REPORT_REVIEW_HIGH_TEMP_BOX 权限",
     "1. 查看待审\n2. 审核操作",
     "—", "功能正常",
     "功能", "否", ""),
    ("E-10", "报告审核", "P1", "振动台报告审核", "REPORT_REVIEW_VIBRATION_TABLE 权限",
     "1. 查看待审\n2. 曲线图预览/下载\n3. 审核",
     "—", "功能正常",
     "功能", "否", ""),
    ("E-11", "报告审核", "P1", "高温高压井报告审核", "REPORT_REVIEW_HIGH_TEMP_PRESSURE 权限",
     "1. 查看待审\n2. 审核操作",
     "—", "功能正常",
     "功能", "否", ""),
    ("E-12", "报告审核", "P1", "模拟井报告审核", "REPORT_REVIEW_SIMULATED_WELL 权限",
     "1. 查看待审\n2. 审核操作",
     "—", "功能正常",
     "功能", "否", ""),
    ("E-13", "测试报告", "P1", "缺陷报告生成", "有缺陷板卡",
     "1. 进入缺陷报告 Tab\n2. 选择板卡\n3. 生成报告",
     "—", "缺陷报告生成成功",
     "功能", "否", ""),
    ("E-14", "测试报告", "P1", "批次报告生成", "有批次数据",
     "1. 进入批次报告 Tab\n2. 选择批次\n3. 生成",
     "—", "批次报告生成成功",
     "功能", "否", ""),
    ("E-15", "疑问工单", "P1", "疑问工单处理闭环", "有疑问工单",
     "1. 查看疑问列表\n2. 打开详情\n3. 直接处理\n4. 验证闭环",
     "—", "处理成功；状态正确更新",
     "功能", "否", ""),

    # ===== P1 数据归档与标准库 =====
    ("F-01", "数据归档", "P1", "历史报告查询归档", "有历史报告",
     "1. 设置查询条件\n2. 查看报告\n3. 执行归档",
     "—", "查询准确；归档操作成功",
     "功能", "否", ""),
    ("F-02", "标准库", "P1", "检测项库 CRUD", "有目录数据",
     "1. 新增一级目录\n2. 新增二级目录\n3. 新增检测项\n4. 编辑删除",
     "检测项: 电压测试",
     "树形结构正确；CRUD 正常",
     "功能", "否", ""),
    ("F-03", "标准库", "P1", "缺陷项库 CRUD", "—",
     "1. 目录树操作\n2. 缺陷项增删改",
     "缺陷项: 焊点虚焊",
     "CRUD 正常",
     "功能", "否", ""),
    ("F-04", "标准库", "P1", "疑问项库 CRUD", "—",
     "1. 目录树操作\n2. 疑问项增删改",
     "—", "CRUD 正常",
     "功能", "否", ""),

    # ===== P1 委托单 =====
    ("G-01", "委托单", "P1", "新增委托单", "—",
     "1. 填写委托单位/联系人\n2. 添加多样品\n3. 关联检测项目\n4. 提交",
     "委托单位: 测试公司A；样品≥2",
     "委托单创建成功",
     "功能", "否", ""),
    ("G-02", "委托单", "P1", "板卡关联汇总", "有委托单和板卡数据",
     "1. 查看已关联/未关联统计\n2. 筛选\n3. 查看绑定关系",
     "—", "统计数据准确",
     "功能", "否", ""),
    ("G-03", "委托单", "P1", "委托单数据管理", "有委托单",
     "1. 查询\n2. 查看详情\n3. 编辑\n4. 删除",
     "—", "CRUD 正常；样品/项目可维护",
     "功能", "否", ""),

    # ===== P1 巡检 =====
    ("H-01", "巡检", "P1", "巡检项配置", "有设备 SN",
     "1. 选择设备 SN\n2. 新增巡检项\n3. 设置是否拍照\n4. 保存",
     "设备 SN: HALT-001",
     "巡检清单保存成功",
     "功能", "否", ""),
    ("H-02", "巡检", "P1", "巡检日历查看", "有巡检记录",
     "1. 切换设备类型\n2. 查看日历\n3. 点击日期查看详情",
     "6类设备: HALT-HASS/水循环/高温高压井/高温箱/振动台/模拟井",
     "日历展示正确；详情完整",
     "功能", "否", ""),

    # ===== P1 权限与签名 =====
    ("I-01", "用户管理", "P1", "用户 CRUD", "USER_MANAGE 权限",
     "1. 新增用户\n2. 分配角色\n3. 编辑简介\n4. 查询",
     "用户名: test_new_user",
     "用户创建成功；角色生效",
     "功能", "否", ""),
    ("I-02", "签名管理", "P1", "签名上传与管理", "SIGNATURE_MANAGE 权限",
     "1. 上传签名图片\n2. 预览\n3. 删除\n4. 验证报告签名显示",
     "PNG/JPG 签名图",
     "签名管理正常；报告嵌入签名",
     "功能", "否", ""),
    ("I-03", "工时管理", "P1", "用户工时查看", "role15 权限",
     "1. 查询用户列表\n2. 选中用户\n3. 查看工时记录",
     "—", "工时数据正确展示",
     "功能", "否", ""),
    ("I-04", "签名", "P1", "我的签名", "已登录",
     "1. 进入 /signature/my\n2. 上传/更新签名\n3. 删除",
     "—", "个人签名管理正常",
     "功能", "否", ""),

    # ===== P1 模拟井 =====
    ("J-01", "模拟井", "P1", "报告模板管理", "—",
     "1. 上传 docx 模板\n2. 解析字段\n3. 配置测试产品/仪器/项目/标准\n4. CRUD",
     "docx 模板文件",
     "模板解析正确；配置保存成功",
     "功能", "否", ""),
    ("J-02", "模拟井", "P1", "模拟井仪器管理", "—",
     "1. 仪器型号 CRUD\n2. 仪器列表 CRUD\n3. 状态管理(正常/到期/维保)",
     "—", "状态流转正确",
     "功能", "否", ""),

    # ===== P2 边界与专项 =====
    ("P2-01", "并发", "P2", "双人同时审核工艺", "2个审核员账号；1条待审工艺",
     "1. 两人同时打开同一工艺\n2. 一人先通过\n3. 另一人后操作",
     "—", "后者提示已处理；无重复审核",
     "边界", "否", ""),
    ("P2-02", "并发", "P2", "双人同时审核报告", "同上",
     "1. 两人同时审核同一报告\n2. 验证互斥",
     "—", "无重复审核",
     "边界", "否", ""),
    ("P2-03", "性能", "P2", "大数据量列表分页", "1000+条工单/报告",
     "1. 打开列表\n2. 翻页\n3. 条件筛选",
     "—", "响应时间可接受(<3s)；分页正确",
     "性能", "否", ""),
    ("P2-04", "异常", "P2", "网络超时处理", "可模拟断网",
     "1. 断开网络\n2. 执行操作\n3. 恢复网络",
     "—", "有明确错误提示；恢复后可重试",
     "边界", "否", ""),
    ("P2-05", "文件", "P2", "超大脚本上传", "有工艺编辑权限",
     "1. 上传超大脚本文件\n2. 上传空文件\n3. 上传非法格式",
     ">10MB 文件；.exe 文件",
     "有大小/格式限制提示",
     "边界", "否", ""),
    ("P2-06", "表单", "P2", "必填项校验", "—",
     "1. 各表单留空提交\n2. 输入非法格式",
     "非法 SN 格式",
     "校验提示准确",
     "边界", "否", ""),
    ("P2-07", "兼容", "P2", "浏览器兼容", "Chrome/Edge 最新版",
     "1. 两浏览器执行 P0 冒烟",
     "Chrome + Edge",
     "功能一致无异常",
     "兼容", "否", ""),
    ("P2-08", "UI", "P2", "主题与布局", "—",
     "1. 切换暗色主题\n2. 折叠侧边栏\n3. 移动端宽度",
     "—", "布局自适应正常",
     "UI", "否", ""),

    # ===== P3 已知限制 =====
    ("P3-01", "数据看板", "P3", "演示数据标注", "—",
     "1. 打开 /data-board\n2. 确认数据来源",
     "—", "演示数据不影响业务判断；已知限制",
     "验收", "否", "占位/演示"),
    ("P3-02", "测试监控", "P3", "进度监控页状态", "—",
     "1. 访问 /testmanage/testmonitor",
     "—", "确认占位或演示状态",
     "验收", "否", "hidden 路由"),
    ("P3-03", "日常管理", "P3", "液氮泄漏监测占位", "—",
     "1. 访问 /exceptionmanage/l-nitrogen",
     "—", "确认为占位页",
     "验收", "否", "hidden 菜单"),
]

# 测试账号矩阵
ACCOUNTS = [
    ("超级管理员", "admin", "全部权限", "全菜单、全操作验证"),
    ("测试工单员", "order_user", "role 11", "工单创建/发布/删除"),
    ("工艺创建员", "process_creator", "TEST_PROCESS_CREATE", "工艺编辑提交"),
    ("工艺审核员", "process_reviewer", "TEST_PROCESS_REVIEW", "工艺审核通过/打回"),
    ("报告审核员", "report_reviewer", "role 16 + REPORT_REVIEW_*", "各类型报告审核"),
    ("板型编辑员", "board_editor", "role 7", "电路板型号新增/编辑"),
    ("测试类型编辑", "type_editor", "role 14", "测试类型维护"),
    ("工时管理员", "worktime_admin", "role 15", "工时管理"),
    ("用户管理员", "user_admin", "role 18 / USER_MANAGE", "用户 CRUD"),
    ("角色管理员", "role_admin", "role 19 / ROLE_MANAGE", "角色权限分配"),
    ("签名管理员", "sign_admin", "SIGNATURE_MANAGE", "签名管理"),
    ("普通只读", "readonly_user", "无特殊权限", "权限拦截验证"),
]

HEADERS = [
    "用例编号", "模块", "优先级", "用例名称", "前置条件",
    "测试步骤", "测试数据", "预期结果", "测试类型", "是否自动化", "备注",
    "执行结果", "执行人", "执行日期", "缺陷编号",
]


def style_header(ws):
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 32


def priority_fill(priority: str):
    colors = {"P0": "FF6B6B", "P1": "4ECDC4", "P2": "FFE66D", "P3": "C8C8C8"}
    return PatternFill("solid", fgColor=colors.get(priority, "FFFFFF"))


def generate_excel():
    wb = Workbook()

    # Sheet1: 测试用例
    ws = wb.active
    ws.title = "测试用例"
    style_header(ws)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    for row_idx, case in enumerate(TEST_CASES, 2):
        for col_idx, val in enumerate(case, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = wrap
            cell.border = border
        prio_cell = ws.cell(row=row_idx, column=3)
        prio_cell.fill = priority_fill(str(prio_cell.value))

    col_widths = [10, 12, 8, 22, 28, 40, 24, 32, 10, 10, 18, 10, 10, 12, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(TEST_CASES) + 1}"

    # Sheet2: 测试账号矩阵
    ws2 = wb.create_sheet("测试账号矩阵")
    acc_headers = ["角色名称", "建议账号", "权限标识", "用途说明"]
    for col, h in enumerate(acc_headers, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row_idx, acc in enumerate(ACCOUNTS, 2):
        for col_idx, val in enumerate(acc, 1):
            ws2.cell(row=row_idx, column=col_idx, value=val).alignment = wrap
    for i, w in enumerate([16, 16, 36, 36], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Sheet3: 功能模块清单
    ws3 = wb.create_sheet("功能模块清单")
    mod_headers = ["一级模块", "二级模块", "路由路径", "主要功能", "权限要求", "状态"]
    modules = [
        ("首页", "待处理任务中心", "/dashboard", "待办聚合、轮询、跳转", "登录即可", "正式"),
        ("首页", "数据看板", "/data-board", "统计卡片、ECharts", "登录即可", "部分演示"),
        ("认证", "登录", "/login", "账号密码、SSO", "公开", "正式"),
        ("认证", "我的签名", "/signature/my", "个人签名管理", "登录即可", "正式"),
        ("装备管理", "电路板管理", "/devicemanage/device/boards", "树形浏览、二维码", "无", "正式"),
        ("装备管理", "仪器管理", "/devicemanage/device/instrument", "CRUD", "无", "正式"),
        ("装备管理", "夹具管理", "/devicemanage/device/fixtures", "CRUD", "无", "正式"),
        ("装备管理", "主机管理", "/devicemanage/device/psi", "CRUD", "无", "正式"),
        ("装备管理", "其他设备管理", "/devicemanage/device/equipments", "CRUD、导入、年检", "无", "正式"),
        ("装备型号", "电路板型号", "/devicemanage/devicemodel/boardmodel", "CRUD", "role 7", "正式"),
        ("装备型号", "电路板组", "/devicemanage/devicemodel/boardgroup", "组内板型配置", "无", "正式"),
        ("装备型号", "其他型号", "/devicemanage/devicemodel/*", "夹具/主机/仪器/产品/其他", "无", "正式"),
        ("测试管理", "测试工单", "/testmanage/testorder", "多Tab工单CRUD、发布、导入导出", "role11/权限码", "正式"),
        ("测试管理", "测试工艺", "/testmanage/testprocess", "工艺CRUD、版本、审核", "创建需权限码", "正式"),
        ("测试管理", "工艺审核", "/testmanage/check", "审核通过/打回", "TEST_PROCESS_REVIEW", "正式"),
        ("测试管理", "测试类型", "/testmanage/testTypelist", "类型维护", "编辑需 role14", "正式"),
        ("测试管理", "疑问工单", "/testmanage/Questionlist", "查询、处理", "无", "正式"),
        ("测试管理", "测试缺陷", "/testmanage/Faultlist", "故障工单", "无", "部分演示"),
        ("测试管理", "测试进度监控", "/testmanage/testmonitor", "进度展示", "hidden", "演示"),
        ("数据管理", "测试报告", "/datamanage/report", "查询、PDF、缺陷/批次报告", "无", "正式"),
        ("数据管理", "报告审核", "/datamanage/reportCheckList", "多类型审核", "role16/权限码", "正式"),
        ("数据管理", "数据归档", "/datamanage/datastore", "历史报告、标准库", "无", "正式"),
        ("数据管理", "工时管理", "/datamanage/work-time", "用户工时", "role 15", "正式"),
        ("数据管理", "委托单", "/datamanage/ordermanage/*", "新增、汇总、管理", "无", "正式"),
        ("数据管理", "模拟井", "/datamanage/simulated-well-data-manage/*", "模板、仪器、审核", "审核需权限码", "正式"),
        ("巡检", "巡检查看", "/inspectionmanage/view", "日历视图", "无", "正式"),
        ("巡检", "巡检添加", "/inspectionmanage/base-info", "巡检项配置", "无", "正式"),
        ("权限", "用户管理", "/permission/user", "用户CRUD", "role18/USER_MANAGE", "正式"),
        ("权限", "角色管理", "/permission/role", "角色权限分配", "role19/ROLE_MANAGE", "正式"),
        ("权限", "签名管理", "/permission/signature", "签名管理", "SIGNATURE_MANAGE", "正式"),
        ("日常管理", "液氮监测", "/exceptionmanage/l-nitrogen", "—", "hidden", "占位"),
    ]
    for col, h in enumerate(mod_headers, 1):
        c = ws3.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for row_idx, mod in enumerate(modules, 2):
        for col_idx, val in enumerate(mod, 1):
            ws3.cell(row=row_idx, column=col_idx, value=val).alignment = wrap
    for i, w in enumerate([14, 16, 42, 36, 24, 10], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # Sheet4: 冒烟检查清单
    ws4 = wb.create_sheet("冒烟检查清单")
    smoke_headers = ["序号", "检查项", "通过标准", "结果(√/×)", "备注"]
    smoke = [
        (1, "登录功能正常", "正确账号可登录，错误账号被拒绝", "", ""),
        (2, "权限路由正确", "各角色菜单与权限一致", "", ""),
        (3, "待办中心可用", "待办展示并跳转正确", "", ""),
        (4, "工艺全流程", "创建→审核→激活", "", ""),
        (5, "工单发布执行", "创建→发布→客户端执行", "", ""),
        (6, "报告审核通过", "审核通过，PDF可下载", "", ""),
        (7, "报告审核打回", "打回后状态正确", "", ""),
    ]
    for col, h in enumerate(smoke_headers, 1):
        c = ws4.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for row_idx, item in enumerate(smoke, 2):
        for col_idx, val in enumerate(item, 1):
            ws4.cell(row=row_idx, column=col_idx, value=val)
    for i, w in enumerate([8, 24, 36, 12, 20], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    out = OUTPUT_DIR / "COSL-test-cases.xlsx"
    wb.save(out)
    print(f"Excel: {out}")
    return out


def new_id():
    return str(uuid.uuid4())


def make_topic(title: str, children: list | None = None) -> dict:
    topic = {"id": new_id(), "class": "topic", "title": title}
    if children:
        topic["children"] = {"attached": children}
    return topic


def build_topic_tree(node) -> dict:
    """将 (标题, [子节点]) 或 字符串 递归转换为 XMind topic"""
    if isinstance(node, str):
        return make_topic(node)
    title, children = node
    return make_topic(title, [build_topic_tree(c) for c in children])


def count_leaves(node) -> int:
    """统计叶子节点(功能测试项)数量"""
    if isinstance(node, str):
        return 1
    _, children = node
    return sum(count_leaves(c) for c in children)


def generate_xmind():
    """生成 XMind Zen 格式 (.xmind)"""
    root_children = [build_topic_tree(branch) for branch in ROOT_CHILDREN]
    leaf_count = sum(count_leaves(b) for b in ROOT_CHILDREN)

    sheet = {
        "id": new_id(),
        "class": "sheet",
        "title": "COSL测试方案",
        "rootTopic": {
            "id": new_id(),
            "class": "topic",
            "title": "COSL测试管理系统\n功能梳理与测试方案",
            "structureClass": "org.xmind.ui.map.unbalanced",
            "children": {"attached": root_children},
        },
    }

    content = [sheet]
    metadata = {
        "creator": {"name": "COSL Test Plan Generator", "version": "1.0"},
        "activeSheetId": sheet["id"],
    }
    manifest = {
        "file-entries": {
            "content.json": {},
            "metadata.json": {},
        }
    }

    out = OUTPUT_DIR / "COSL-test-plan.xmind"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("content.json", json.dumps(content, ensure_ascii=False, indent=2))
        zf.writestr("metadata.json", json.dumps(metadata, ensure_ascii=False, indent=2))
        zf.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    print(f"XMind: {out} ({leaf_count} 个节点)")
    return out


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    generate_excel()
    generate_xmind()
    print(f"\n已生成 {len(TEST_CASES)} 条测试用例")
    print(f"输出目录: {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
